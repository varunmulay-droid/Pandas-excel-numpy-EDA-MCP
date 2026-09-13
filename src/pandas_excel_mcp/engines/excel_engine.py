"""openpyxl-backed Excel read/write/format/chart operations.

Works on an in-memory ``openpyxl.Workbook`` loaded from a dataset's
persisted file, or a freshly created workbook for export scenarios.
"""
from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.chart import AreaChart, BarChart, LineChart, PieChart, ScatterChart, Reference, Series
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import range_boundaries
from openpyxl.utils.cell import coordinate_from_string

from ..schemas.exceptions import OperationError, SheetNotFoundError, ValidationError

_CHART_TYPES = {"bar": BarChart, "line": LineChart, "pie": PieChart, "scatter": ScatterChart, "area": AreaChart}


def load_workbook_bytes(raw_bytes: bytes) -> Workbook:
    return load_workbook(io.BytesIO(raw_bytes), data_only=False)


def workbook_to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def list_sheets(wb: Workbook) -> list[dict[str, Any]]:
    return [
        {"name": ws.title, "dimensions": ws.dimensions, "max_row": ws.max_row, "max_column": ws.max_column}
        for ws in wb.worksheets
    ]


def _get_sheet(wb: Workbook, sheet_name: str):
    if sheet_name not in wb.sheetnames:
        raise SheetNotFoundError(
            f"Sheet '{sheet_name}' was not found.", {"available_sheets": wb.sheetnames}
        )
    return wb[sheet_name]


def read_range(wb: Workbook, sheet_name: str, cell_range: str) -> list[list[Any]]:
    ws = _get_sheet(wb, sheet_name)
    try:
        cells = ws[cell_range]
    except Exception as exc:  # noqa: BLE001
        raise ValidationError(f"Invalid range '{cell_range}': {exc}") from exc
    if not isinstance(cells, tuple):
        cells = ((cells,),)
    return [[c.value for c in row] for row in cells]


def write_range(
    wb: Workbook, sheet_name: str, start_cell: str, values: list[list[Any]]
) -> dict[str, Any]:
    ws = _get_sheet(wb, sheet_name) if sheet_name in wb.sheetnames else wb.create_sheet(sheet_name)
    col0, row0 = coordinate_from_string(start_cell)
    from openpyxl.utils import column_index_from_string

    col0_idx = column_index_from_string(col0)
    for r, row_vals in enumerate(values):
        for c, val in enumerate(row_vals):
            ws.cell(row=row0 + r, column=col0_idx + c, value=val)
    return {"sheet": sheet_name, "written_rows": len(values), "start_cell": start_cell}


def excel_format(
    wb: Workbook,
    sheet_name: str,
    cell_range: str,
    font: dict[str, Any] | None = None,
    fill_color: str | None = None,
    border: bool = False,
    alignment: str | None = None,
    number_format: str | None = None,
    column_widths: dict[str, float] | None = None,
    row_heights: dict[str, float] | None = None,
    freeze_panes: str | None = None,
) -> dict[str, Any]:
    ws = _get_sheet(wb, sheet_name)
    cells = ws[cell_range]
    if not isinstance(cells, tuple):
        cells = ((cells,),)

    font_obj = Font(**font) if font else None
    fill_obj = (
        PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        if fill_color
        else None
    )
    border_obj = (
        Border(*(Side(style="thin"),) * 4) if border else None
    )
    align_obj = Alignment(horizontal=alignment) if alignment else None

    for row in cells:
        for cell in row:
            if font_obj:
                cell.font = font_obj
            if fill_obj:
                cell.fill = fill_obj
            if border_obj:
                cell.border = border_obj
            if align_obj:
                cell.alignment = align_obj
            if number_format:
                cell.number_format = number_format

    for col, width in (column_widths or {}).items():
        ws.column_dimensions[col].width = width
    for row_num, height in (row_heights or {}).items():
        ws.row_dimensions[int(row_num)].height = height
    if freeze_panes:
        ws.freeze_panes = freeze_panes

    return {"sheet": sheet_name, "range": cell_range, "formatted": True}


def excel_chart(
    wb: Workbook,
    sheet_name: str,
    chart_type: str,
    data_range: str,
    categories_range: str | None,
    anchor: str,
    title: str | None = None,
) -> dict[str, Any]:
    if chart_type not in _CHART_TYPES:
        raise ValidationError(f"Unsupported chart type '{chart_type}'.", {"allowed": list(_CHART_TYPES)})
    ws = _get_sheet(wb, sheet_name)
    chart = _CHART_TYPES[chart_type]()
    if title:
        chart.title = title

    min_col, min_row, max_col, max_row = range_boundaries(data_range.split("!")[-1])
    data_ref = Reference(ws, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)
    chart.add_data(data_ref, titles_from_data=True)

    if categories_range:
        c_min_col, c_min_row, c_max_col, c_max_row = range_boundaries(categories_range.split("!")[-1])
        cats_ref = Reference(ws, min_col=c_min_col, min_row=c_min_row, max_col=c_max_col, max_row=c_max_row)
        chart.set_categories(cats_ref)

    ws.add_chart(chart, anchor)
    return {"sheet": sheet_name, "chart_type": chart_type, "anchor": anchor}
