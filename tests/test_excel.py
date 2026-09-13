import io

import pandas as pd
import pytest
from openpyxl import Workbook

from src.pandas_excel_mcp.engines import excel_engine
from src.pandas_excel_mcp.schemas.exceptions import SheetNotFoundError


def _make_workbook_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Region", "Sales"])
    ws.append(["West", 12000])
    ws.append(["East", 9500])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_list_sheets():
    wb = excel_engine.load_workbook_bytes(_make_workbook_bytes())
    sheets = excel_engine.list_sheets(wb)
    assert sheets[0]["name"] == "Sheet1"


def test_read_range():
    wb = excel_engine.load_workbook_bytes(_make_workbook_bytes())
    values = excel_engine.read_range(wb, "Sheet1", "A1:B3")
    assert values[0] == ["Region", "Sales"]
    assert values[1] == ["West", 12000]


def test_read_range_missing_sheet_raises():
    wb = excel_engine.load_workbook_bytes(_make_workbook_bytes())
    with pytest.raises(SheetNotFoundError):
        excel_engine.read_range(wb, "NoSuchSheet", "A1:B2")


def test_write_range():
    wb = excel_engine.load_workbook_bytes(_make_workbook_bytes())
    result = excel_engine.write_range(wb, "Sheet1", "D1", [["New", "Data"], [1, 2]])
    assert result["written_rows"] == 2
    values = excel_engine.read_range(wb, "Sheet1", "D1:E2")
    assert values[0] == ["New", "Data"]


def test_excel_format_applies_fill():
    wb = excel_engine.load_workbook_bytes(_make_workbook_bytes())
    result = excel_engine.excel_format(wb, "Sheet1", "A1:B1", fill_color="FFFF00")
    assert result["formatted"] is True


def test_excel_chart_bar():
    wb = excel_engine.load_workbook_bytes(_make_workbook_bytes())
    result = excel_engine.excel_chart(
        wb, "Sheet1", "bar", "B1:B3", None, "D5", title="Sales Chart"
    )
    assert result["chart_type"] == "bar"


def test_workbook_roundtrip_bytes():
    raw = _make_workbook_bytes()
    wb = excel_engine.load_workbook_bytes(raw)
    out = excel_engine.workbook_to_bytes(wb)
    assert isinstance(out, bytes) and len(out) > 0
